"""Office-local queue runner. No scheduling or database migrations occur at startup."""
import argparse
import hashlib
import csv
import json
import os
import subprocess
import time
import signal
import threading
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo
from functools import lru_cache

FEEDS = {
    'normal_sales': ('daily_sales', 'sales_date', 'sold_qty'),
    'rtc': ('rtc_sales', 'rtc_date', 'rtc_qty'),
    'wastage': ('wastage', 'wastage_date', 'wastage_qty'),
    'book_stock': ('book_stock', 'book_stock_date', 'book_stock_qty'),
}

CRM_STORE_CODES = {'HNS': '0365', 'PVL': '0766', 'EAS': '0650', 'HYS': '0665',
                   'GRV': '0566', 'WTF': '0602', 'STR': '0388'}


def crm_store_code(code):
    """Translate the application's store identity to its known CRM login marker."""
    import re
    value = str(code).strip()
    if re.fullmatch(r'\d{4}', value):
        return value
    if value.upper() in CRM_STORE_CODES:
        return CRM_STORE_CODES[value.upper()]
    raise ValueError('Unknown CRM store code; configure a verified store identity before downloading')


def enabled_movement_sources(config):
    if config.get('store_settings_source') != 'database':
        return ['normal_sales', 'rtc', 'wastage']
    feeds = config.get('store_defaults', {}).get('feeds', {})
    return [source for source in ('normal_sales', 'rtc', 'wastage')
            if feeds.get(source, {}).get('complete_store_report') is True]


def london_today():
    return datetime.now(ZoneInfo('Europe/London')).date()


def get_store_config(conn, config, store_id):
    if config.get('store_settings_source') != 'database':
        return config['stores'][str(store_id)]
    row = conn.execute('''SELECT s.code,s.erp_type,c.crm_url FROM stores s
        JOIN store_crm_settings c ON c.store_id=s.id
        WHERE s.id=%s AND s.is_active=true''', (store_id,)).fetchone()
    if not row:
        raise ValueError('Active store CRM settings not found; save the URL in CRM Downloads')
    store = {**config['store_defaults'], **config.get('stores', {}).get(str(store_id), {})}
    # Never cache URLs in the worker: operator edits apply to the next claimed job.
    from urllib.parse import urlsplit
    url = urlsplit(row[2])
    if url.scheme not in ('http', 'https') or not url.hostname or url.username or url.password or url.query or url.fragment:
        raise ValueError('Invalid saved store CRM URL')
    return {**store, 'code': row[0], 'erp': row[1], 'crm_url': row[2]}


def nightly_store_ids(conn, config):
    if config.get('store_settings_source') == 'database':
        return [row[0] for row in conn.execute('''SELECT s.id FROM stores s
            JOIN store_crm_settings c ON c.store_id=s.id
            WHERE s.is_active=true AND s.erp_type='msp' AND c.nightly_enabled=true
            ORDER BY s.id''').fetchall()]
    return [int(key) for key in config['stores']]


class ImportValidationError(ValueError):
    def __init__(self, message, details):
        super().__init__(message)
        self.details = details


def prepare_rows(rows, mapping, source, start, end, allow_empty=False):
    """Strict file validation before any writes; repeated product/date lines sum."""
    if source not in FEEDS:
        raise ValueError('Unknown source')
    start, end = date.fromisoformat(start), date.fromisoformat(end)
    if start > end or (source == 'book_stock' and start != end):
        raise ValueError('Invalid report date range')
    result = {}
    for index, row in enumerate(rows, 2):
        code = str(row.get(mapping['code'], '')).strip()
        raw = row.get(mapping['quantity'])
        if not code or raw is None or str(raw).strip() == '':
            raise ValueError(f'Row {index}: missing code or quantity')
        qty = Decimal(str(raw).replace(',', '').strip())
        if not qty.is_finite() or abs(qty) >= Decimal('100000000000000'):
            raise ValueError(f'Row {index}: invalid quantity')
        if qty != qty.quantize(Decimal('0.0001')):
            raise ValueError(f'Row {index}: more than four decimal places')
        value = row.get(mapping.get('date', ''))
        if isinstance(value, datetime):
            value = value.date().isoformat()
        if not value and start != end:
            raise ValueError(f'Row {index}: daily date required for multi-day reports')
        day = date.fromisoformat(str(value or start))
        if not start <= day <= end:
            raise ValueError(f'Row {index}: date outside requested range')
        if source in ('rtc', 'wastage') and qty < 0:
            raise ValueError(f'Row {index}: expected positive movement quantity')
        key = (code, day)
        result[key] = result.get(key, Decimal(0)) + qty
    if not result and not (allow_empty and source != 'book_stock'):
        raise ValueError('Empty report; previous data retained')
    return result


def read_rows(path, mapping):
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            return list(csv.DictReader(handle))
    if path.suffix.lower() != '.xlsx':
        raise ValueError('Expected CSV or XLSX report')
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[mapping['sheet']] if mapping.get('sheet') else book.active
        iterator = sheet.iter_rows(min_row=mapping.get('header_row', 1), values_only=True)
        headers = [str(value).strip() if value is not None else '' for value in next(iterator)]
        return [dict(zip(headers, row)) for row in iterator if any(v is not None for v in row)]
    finally:
        book.close()


def enqueue_nightly(conn, store_ids, day, sources=None):
    """Queue each movement report day once; bounded recovery is scheduled separately."""
    date.fromisoformat(day)
    with conn.transaction():
        for store_id in sorted(set(store_ids)):
            for source in (sources if sources is not None else ('normal_sales', 'rtc', 'wastage')):
                if source not in ('normal_sales', 'rtc', 'wastage'):
                    raise ValueError('Nightly imports support movement feeds only')
                conn.execute('SELECT pg_advisory_xact_lock(73501, %s)', (store_id,))
                conn.execute('''INSERT INTO data_import_runs
                    (store_id, source_type, trigger_type, source_start_date, source_end_date)
                    SELECT %s, %s, 'nightly', %s, %s
                    WHERE EXISTS (SELECT 1 FROM stores WHERE id=%s AND is_active=true)
                    AND NOT EXISTS (SELECT 1 FROM data_import_runs WHERE store_id=%s
                      AND source_type=%s AND source_start_date<=%s AND source_end_date>=%s)''',
                    (store_id, source, day, day, store_id, store_id, source, day, day))


@lru_cache(maxsize=1)
def worker_revision():
    """Fingerprint executable worker/parser code so format failures retry after an upgrade."""
    app = Path(__file__).resolve().parent.parent
    files = [Path(__file__).resolve(), app / 'automation/normalize-report.ts', app / 'automation/nebula_download.py']
    files += [p for p in (app / 'src/parsers').rglob('*.ts') if not p.name.endswith('.test.ts')]
    digest = hashlib.sha256()
    for path in sorted(files):
        if path.is_file():
            digest.update(str(path.relative_to(app)).encode())
            digest.update(path.read_bytes())
    return digest.hexdigest()


def failure_category(message, exception=None):
    value = str(message).lower()
    if isinstance(exception, ImportValidationError) or 'unmatched' in value:
        return 'product_review'
    if any(word in value for word in ('another store', 'wrong store', 'store identity', 'missing store code',
            'login code', 'complete_store_report', 'includes_negative_stock', 'configured store',
            'invalid file extension', 'date differs', 'dates differ', 'date outside', 'row date differs',
            'date controls changed', 'did not retain', 'empty report', 'no records available', 'four decimal',
            'negative movement', 'expected positive', 'invalid quantity', 'missing code or quantity')):
        return 'review_required'
    if 'parsing failed' in value or 'columns not detected' in value:
        return 'parser_upgrade'
    if isinstance(exception, subprocess.TimeoutExpired) or any(word in value for word in (
            'timeout', 'timed out', 'login page is unavailable', 'downloader failed', 'crm report export failed',
            'downloader did not produce', 'connection', 'worker stopped')):
        return 'download_retry'
    return 'review_required'


def enqueue_failed_retries(conn, store_ids, sources, day, revision=None):
    """Retry recent transient failures at most three times; never loop product/parser errors.

    The installed timer invokes this at 02:00, 03:00, 04:00 and 05:00 London time.
    Each candidate must have cooled down for 45 minutes. The last 30 report days
    are eligible; operator controls, active work and newer attempts take precedence.
    """
    revision = revision or worker_revision()
    last_day = date.fromisoformat(day)
    first_day = last_day - timedelta(days=29)
    queued = []
    with conn.transaction():
        for store_id in sorted(set(store_ids)):
            conn.execute('SELECT pg_advisory_xact_lock(73501, %s)', (store_id,))
            failures = conn.execute("""SELECT r.id,r.source_type,r.source_start_date,r.source_end_date,r.details,r.error_message
                FROM data_import_runs r JOIN stores s ON s.id=r.store_id AND s.is_active=true
                WHERE r.store_id=%s AND r.source_type=ANY(%s) AND r.status='failed'
                  AND r.source_start_date BETWEEN %s AND %s AND r.source_end_date<=%s
                  AND COALESCE(r.completed_at,r.started_at)<now()-interval '45 minutes'
                  AND NOT(r.details ? 'queue_control') AND NOT(r.details ? 'unmatched_barcodes')
                  AND NOT(r.details ? 'unknown_product_codes') AND NOT(r.details ? 'automatic_retry_as')
                  AND NOT EXISTS(SELECT 1 FROM data_import_runs n WHERE n.store_id=r.store_id AND n.source_type=r.source_type
                    AND n.source_start_date<=r.source_end_date AND n.source_end_date>=r.source_start_date
                    AND (n.id>r.id OR n.status IN ('queued','downloading','downloaded','importing')))
                ORDER BY r.id LIMIT 100""", (store_id, list(sources), first_day, last_day, last_day)).fetchall()
            for run_id, source, start, end, details, error in failures:
                category = failure_category(error or '')
                count = details.get('automatic_retry_count', 0)
                if type(count) is not int or not 0 <= count < 3:
                    continue
                if category == 'parser_upgrade':
                    if details.get('worker_revision') == revision:
                        continue
                elif category != 'download_retry':
                    continue
                payload = {'retry_of': run_id, 'automatic_retry_count': count + 1,
                           'automatic_retry_reason': category, 'scheduled_worker_revision': revision}
                created = conn.execute("""INSERT INTO data_import_runs(store_id,source_type,trigger_type,
                    source_start_date,source_end_date,details) VALUES(%s,%s,'retry',%s,%s,%s::jsonb) RETURNING id""",
                    (store_id, source, start, end, json.dumps(payload))).fetchone()[0]
                conn.execute("UPDATE data_import_runs SET details=details || %s::jsonb WHERE id=%s",
                             (json.dumps({'automatic_retry_as': created}), run_id))
                queued.append(created)
    return queued


def claim(conn, config=None):
    # Advisory locking during imports serializes writes for a store/source.
    with conn.transaction():
        if config and config.get('store_settings_source') == 'database':
            # This Nebula profile supports the three movement reports only.
            # Leave manual Book Stock requests for a configured stock downloader.
            return conn.execute('''UPDATE data_import_runs SET status='downloading', started_at=now()
                WHERE id=(SELECT r.id FROM data_import_runs r
                  JOIN stores s ON s.id=r.store_id
                  JOIN store_crm_settings c ON c.store_id=s.id
                  WHERE r.status='queued' AND s.is_active=true AND s.erp_type='msp'
                    AND r.source_type = ANY(%s)
                    AND (%s::int IS NULL OR r.store_id=%s)
                  ORDER BY r.source_start_date ASC NULLS LAST,r.id ASC FOR UPDATE OF r SKIP LOCKED LIMIT 1)
                RETURNING id,store_id,source_type,source_start_date,source_end_date''',
                (enabled_movement_sources(config), config.get('only_store_id'), config.get('only_store_id'))).fetchone()
        return conn.execute('''UPDATE data_import_runs SET status='downloading', started_at=now()
            WHERE id=(SELECT id FROM data_import_runs WHERE status='queued'
                ORDER BY source_start_date ASC NULLS LAST,id ASC FOR UPDATE SKIP LOCKED LIMIT 1)
            RETURNING id, store_id, source_type, source_start_date, source_end_date''').fetchone()


def row_fingerprint(rows):
    return hashlib.sha256(json.dumps([(pid, str(day), format(qty.normalize(), 'f'))
        for (pid, day), qty in sorted(rows.items())], separators=(',', ':')).encode()).hexdigest()


def import_report(conn, job, rows, filename, mapping_type="store_product_code", source_sha256=None):
    run_id, store_id, source, start, end = job
    table, date_column, quantity_column = FEEDS[source]  # Identifiers only from allowlist.
    with conn.transaction():
        conn.execute('SELECT pg_advisory_xact_lock(73501, %s)', (store_id,))
        state = conn.execute('SELECT status FROM data_import_runs WHERE id=%s FOR UPDATE', (run_id,)).fetchone()
        if not state or state[0] != 'importing':
            raise ValueError('Job is no longer importing')
        newer = conn.execute("""SELECT id FROM data_import_runs WHERE store_id=%s AND source_type=%s
            AND id>%s AND source_start_date<=%s AND source_end_date>=%s
            AND (status='success' OR details->>'partial_import'='true') LIMIT 1""",
            (store_id, source, run_id, end, start)).fetchone()
        if newer:
            raise ValueError('A newer overlapping report is already saved; previous data retained')
        if mapping_type == 'barcode':
            products = dict(conn.execute('SELECT barcode, store_product_id FROM product_barcodes WHERE store_id=%s', (store_id,)).fetchall())
        elif mapping_type == 'store_product_code':
            products = dict(conn.execute('SELECT identifier,store_product_id FROM crm_product_code_aliases WHERE store_id=%s', (store_id,)).fetchall())
            native = dict(conn.execute('SELECT store_product_code, id FROM store_products WHERE store_id=%s', (store_id,)).fetchall())
            if any(code in native and native[code] != target for code, target in products.items()):
                raise ValueError('Product code alias conflicts with product master; manual review required')
            products.update(native)
        else:
            raise ValueError('Unsupported product mapping')
        unknown = sorted({code for code, _ in rows if code not in products})
        resolved = {}
        for (code, day), quantity in rows.items():
            if code not in products:
                continue
            key = (products[code], day)
            resolved[key] = resolved.get(key, Decimal(0)) + quantity
        if unknown:
            details = {
                'unmatched_barcodes' if mapping_type == 'barcode' else 'unknown_product_codes': unknown,
                'unmatched_identifier_count': len(unknown),
            }
            if source == 'book_stock':
                # A stock snapshot drives every later stock calculation and remains atomic.
                raise ImportValidationError(f'{len(unknown)} unmatched product identifiers; stock snapshot retained', details)
            return publish_partial_report(conn, job, rows, resolved, filename, details, source_sha256)
        rows = resolved
        fingerprint = row_fingerprint(rows)
        details = {'data_sha256': fingerprint, 'source_sha256': source_sha256, 'zero_activity': source in ('wastage', 'rtc') and not rows, 'activity': ('Zero wastage' if source == 'wastage' else 'Zero RTC activity') if source in ('wastage', 'rtc') and not rows else 'Imported'}
        # Only the latest overlapping import describes the current database state.
        # An older matching fingerprint must not suppress a legitimate correction.
        previous = conn.execute("""SELECT id, imported_at, details->>'data_sha256', details->>'duplicate_of'
            FROM data_import_runs WHERE store_id=%s AND source_type=%s
            AND source_start_date<=%s AND source_end_date>=%s AND status='success' AND COALESCE(details->>'has_unmatched_products','false')<>'true'
            ORDER BY imported_at DESC NULLS LAST, id DESC LIMIT 1""",
            (store_id, source, end, start)).fetchone()
        unchanged = previous and previous[2] == fingerprint
        if unchanged and source == 'book_stock':
            # Browser uploads need not have a worker audit row. A newer manual
            # snapshot invalidates the previous worker fingerprint.
            latest = conn.execute("""SELECT max(imported_at) FROM book_stock WHERE
                store_product_id IN (SELECT id FROM store_products WHERE store_id=%s)""", (store_id,)).fetchone()[0]
            unchanged = latest == previous[1]
        elif unchanged:
            # Likewise, manual movement imports can change this range without
            # changing the worker audit history. Compare the stored data too.
            current = conn.execute(f'SELECT store_product_id,{date_column},{quantity_column} FROM {table} '
                f'WHERE {date_column} BETWEEN %s AND %s AND store_product_id IN '
                '(SELECT id FROM store_products WHERE store_id=%s)', (start, end, store_id)).fetchall()
            unchanged = row_fingerprint({(pid, day): qty for pid, day, qty in current}) == fingerprint
        if unchanged:
            details['duplicate_of'] = int(previous[3]) if previous[3] else previous[0]
            conn.execute("""UPDATE data_import_runs SET status='success', imported_at=%s, completed_at=now(),
                imported_row_count=0, rejected_row_count=0, source_file_name=%s, details=details || %s::jsonb WHERE id=%s""",
                (previous[1], filename, json.dumps(details), run_id))
            conn.execute('DELETE FROM crm_partial_publications WHERE store_id=%s AND source_type=%s AND report_date BETWEEN %s AND %s', (store_id, source, start, end))
            conn.execute('DELETE FROM crm_report_stage WHERE run_id=%s', (run_id,))
            return {'status': 'success', 'saved': 0}
        imported_at = conn.execute('SELECT clock_timestamp()').fetchone()[0]
        if source == 'book_stock':
            # One snapshot timestamp, committed atomically with the successful audit record.
            for (product_id, day), quantity in rows.items():
                conn.execute('''INSERT INTO book_stock (store_product_id, book_stock_date, book_stock_qty, imported_at)
                    VALUES (%s,%s,%s,%s)''', (product_id, day, quantity, imported_at))
            conn.execute('DELETE FROM not_posted_deliveries WHERE store_product_id IN (SELECT id FROM store_products WHERE store_id=%s)', (store_id,))
        else:
            # Reports must cover the whole store/day. Replace the requested range to remove
            # corrected/deleted lines on retry, without touching the other two feeds.
            conn.execute(f'DELETE FROM {table} WHERE {date_column} BETWEEN %s AND %s AND store_product_id IN (SELECT id FROM store_products WHERE store_id=%s)', (start, end, store_id))
            for (product_id, day), quantity in rows.items():
                conn.execute(f'INSERT INTO {table} (store_product_id,{date_column},{quantity_column}) VALUES (%s,%s,%s)', (product_id, day, quantity))
        conn.execute('''UPDATE data_import_runs SET status='success', imported_at=%s, completed_at=now(),
            imported_row_count=%s, rejected_row_count=0, source_file_name=%s, details=details || %s::jsonb WHERE id=%s''',
            (imported_at, len(rows), filename, json.dumps(details), run_id))
        conn.execute('DELETE FROM crm_partial_publications WHERE store_id=%s AND source_type=%s AND report_date BETWEEN %s AND %s', (store_id, source, start, end))
        conn.execute('DELETE FROM crm_report_stage WHERE run_id=%s', (run_id,))
        return {'status': 'success', 'saved': len(rows)}


def publish_partial_report(conn, job, original, resolved, filename, details, source_sha256):
    """Publish known subtotals only into empty or verified partial-owned store days.

    A complete previous import or a manual edit is never replaced by a partial
    report. Replays replace the owned day from the full stage, never add quantities.
    This function runs inside import_report's store-locked transaction.
    """
    run_id, store_id, source, start, end = job
    table, date_column, quantity_column = FEEDS[source]
    published, protected = [], []
    saved = 0
    for day in sorted({day for _, day in original}):
        matched = {key: qty for key, qty in resolved.items() if key[1] == day}
        if not matched:
            continue
        current = conn.execute(f'SELECT store_product_id,{date_column},{quantity_column} FROM {table} '
            f'WHERE {date_column}=%s AND store_product_id IN (SELECT id FROM store_products WHERE store_id=%s)',
            (day, store_id)).fetchall()
        owner = conn.execute('SELECT run_id,data_sha256 FROM crm_partial_publications WHERE store_id=%s AND source_type=%s AND report_date=%s FOR UPDATE',
                             (store_id, source, day)).fetchone()
        complete = conn.execute("""SELECT id FROM data_import_runs WHERE store_id=%s AND source_type=%s
            AND status='success' AND COALESCE(details->>'has_unmatched_products','false')<>'true' AND source_start_date<=%s AND source_end_date>=%s LIMIT 1""",
            (store_id, source, day, day)).fetchone()
        current_hash = row_fingerprint({(pid, d): qty for pid, d, qty in current})
        if complete or (current and (not owner or owner[1] != current_hash)) or (owner and owner[0] > run_id):
            protected.append(str(day))
            continue
        conn.execute(f'DELETE FROM {table} WHERE {date_column}=%s AND store_product_id IN (SELECT id FROM store_products WHERE store_id=%s)', (day, store_id))
        for (pid, d), qty in matched.items():
            conn.execute(f'INSERT INTO {table}(store_product_id,{date_column},{quantity_column}) VALUES(%s,%s,%s)', (pid, d, qty))
        conn.execute("""INSERT INTO crm_partial_publications(store_id,source_type,report_date,run_id,data_sha256,published_row_count)
            VALUES(%s,%s,%s,%s,%s,%s) ON CONFLICT(store_id,source_type,report_date)
            DO UPDATE SET run_id=EXCLUDED.run_id,data_sha256=EXCLUDED.data_sha256,
            published_row_count=EXCLUDED.published_row_count,updated_at=now()""",
            (store_id, source, day, run_id, row_fingerprint(matched), len(matched)))
        published.append(str(day))
        saved += len(matched)
    details.update({'has_unmatched_products': True, 'partial_import': bool(published), 'partial_import_version': 1,
                    'partial_published_dates': published, 'protected_complete_dates': protected,
                    'matched_staged_rows': len(resolved), 'source_sha256': source_sha256})
    unknown_set = set(details.get('unmatched_barcodes', details.get('unknown_product_codes', [])))
    # Persist review items in the same transaction as matching quantities.
    kind = 'Barcode' if 'unmatched_barcodes' in details else 'Product code'
    conn.execute("""INSERT INTO crm_product_identifiers(store_id,kind,identifier)
        SELECT %s,%s,unnest(%s::text[]) ON CONFLICT(store_id,kind,identifier) DO NOTHING""",
        (store_id, kind, sorted(unknown_set)))
    conn.execute("""INSERT INTO crm_identifier_occurrences(identifier_id,run_id)
        SELECT id,%s FROM crm_product_identifiers WHERE store_id=%s AND kind=%s AND identifier=ANY(%s::text[])
        ON CONFLICT(identifier_id,run_id) DO NOTHING""", (run_id,store_id,kind,sorted(unknown_set)))
    pending_rows = sum(1 for code, _ in original if code in unknown_set)
    message = f'{saved} matched product/date rows saved; {len(details.get("unmatched_barcodes", details.get("unknown_product_codes", [])))} unmatched identifiers moved to Unmatched Products. Import completed with review items.'
    if protected:
        message += ' Earlier complete or manually edited data retained for protected dates; matched rows remain staged there.'
    conn.execute("""UPDATE data_import_runs SET status='success',completed_at=now(),
        imported_at=CASE WHEN %s>0 THEN clock_timestamp() ELSE NULL END,
        imported_row_count=%s,rejected_row_count=%s,source_file_name=%s,error_message=NULL,
        details=details || %s::jsonb WHERE id=%s""",
        (saved, saved, pending_rows, filename, json.dumps({**details, "review_message": message}), run_id))
    return {'status': 'partial' if published else 'pending', 'saved': saved, 'details': details, 'message': message}


def stage_report(conn, job, rows, mapping_type, source_sha256):
    """Retain the full validated input for progressive matching and exact finalization."""
    run_id, store_id, source, start, end = job
    with conn.transaction():
        conn.execute('SELECT pg_advisory_xact_lock(73501, %s)', (store_id,))
        state = conn.execute('SELECT status FROM data_import_runs WHERE id=%s FOR UPDATE', (run_id,)).fetchone()
        if not state or state[0] not in ('downloaded', 'downloading'):
            raise ValueError('Job is no longer available for staging')
        conn.execute('DELETE FROM crm_report_stage WHERE run_id=%s', (run_id,))
        # Batches avoid one database round trip for every transaction in a large report.
        values = [(run_id, code, str(day), str(qty), mapping_type) for (code, day), qty in rows.items()]
        for offset in range(0, len(values), 1000):
            batch = values[offset:offset + 1000]
            conn.execute('INSERT INTO crm_report_stage(run_id,identifier,report_date,quantity,mapping_type) VALUES '
                         + ','.join(['(%s,%s,%s,%s,%s)'] * len(batch)), tuple(v for row in batch for v in row))
        details = {'staged_data_sha256': row_fingerprint(rows), 'staged_row_count': len(rows),
                   'source_sha256': source_sha256, 'staged_mapping_type': mapping_type}
        conn.execute("UPDATE data_import_runs SET status='importing',downloaded_row_count=%s,details=details || %s::jsonb WHERE id=%s",
                     (len(rows), json.dumps(details), run_id))


def reuse_validated_stage(conn, job):
    """Only reuse an immutable, hashed, validated parent for the same store/feed/dates."""
    run_id, store_id, source, start, end = job
    parent = conn.execute("""SELECT p.id,p.details,p.source_file_name,p.source_file_path FROM data_import_runs r
        JOIN data_import_runs p ON p.id=CASE WHEN (r.details->>'retry_of') ~ '^[0-9]+$'
            THEN (r.details->>'retry_of')::bigint ELSE NULL END
        WHERE r.id=%s AND r.details->>'reuse_validated_stage'='true' AND (p.status='failed' OR p.details->>'has_unmatched_products'='true')
          AND p.store_id=%s AND p.source_type=%s AND p.source_start_date=%s AND p.source_end_date=%s""",
        (run_id, store_id, source, start, end)).fetchone()
    if not parent:
        return False
    parent_id, details, filename, filepath = parent
    if not details.get('staged_row_count') or not details.get('source_sha256'):
        return False
    staged = conn.execute('SELECT identifier,report_date,quantity,mapping_type FROM crm_report_stage WHERE run_id=%s', (parent_id,)).fetchall()
    if len(staged) != details['staged_row_count'] or any(row[3] != details.get('staged_mapping_type') for row in staged):
        return False
    rows = {(code, day): qty for code, day, qty, _ in staged}
    if row_fingerprint(rows) != details.get('staged_data_sha256'):
        return False
    stage_report(conn, job, rows, details['staged_mapping_type'], details['source_sha256'])
    conn.execute("UPDATE data_import_runs SET source_file_path=%s,downloaded_at=now(),details=details || %s::jsonb WHERE id=%s AND status='importing'",
                 (filepath, json.dumps({'reused_stage_of': parent_id}), run_id))
    result = import_report(conn, job, rows, filename or 'retained-report', details['staged_mapping_type'], details['source_sha256'])
    return result


def process_job(conn, config, job):
    run_id, store_id, source, start, end = job
    review_descriptions = {}
    try:
        conn.execute("UPDATE data_import_runs SET details=details || %s::jsonb WHERE id=%s AND status NOT IN ('success','failed')",
                     (json.dumps({'worker_revision': worker_revision()}), run_id))
        store = get_store_config(conn, config, store_id)
        identity = conn.execute('SELECT code FROM stores WHERE id=%s AND is_active=true', (store_id,)).fetchone()
        if not identity or (store.get('code') and store['code'] != identity[0]):
            raise ValueError('Configured store identity does not match the database')
        feed = store['feeds'][source]
        if source != 'book_stock' and feed.get('complete_store_report') is not True:
            raise ValueError('Movement imports require complete_store_report=true')
        if source == 'book_stock' and feed.get('includes_negative_stock') is not True:
            raise ValueError('Book Stock report must include negative stock')
        start = str(start or london_today())
        end = str(end or start)
        reused = reuse_validated_stage(conn, (run_id, store_id, source, start, end))
        if reused:
            print(f'Job {run_id}: {reused["status"]} (retained validated report) ({reused["saved"]} rows saved)', flush=True)
            return
        folder = Path(config['download_directory']).resolve() / str(run_id)
        folder.mkdir(parents=True, exist_ok=False)  # Never reuse an old download.
        extension = feed.get('extension', '.xlsx')
        if extension not in ('.csv', '.xlsx', '.xls') or (extension == '.xls' and not feed.get('normalize_with_app_parsers')):
            raise ValueError('Invalid file extension')
        output = folder / ('report' + extension)
        values = dict(store_id=str(store_id), store_code=crm_store_code(identity[0]) if config.get('store_settings_source') == 'database' else identity[0], source=source, start=start, end=end, output=str(output), crm_url=store.get('crm_url', ''))
        command = [part.format(**values) for part in feed['command']]
        downloaded = subprocess.run(command, check=False, timeout=feed.get('timeout_seconds', 2400), shell=False,
                                    stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
        if downloaded.returncode:
            safe_errors = {
                'wrong_store': 'CRM displayed another or missing store code. Check the saved store URL before retrying.',
                'date_controls': 'CRM daily date controls changed. Verify Date Range, From Date and To Date before retrying.',
                'login_config': 'CRM login code is missing from the server environment.',
                'login_page': 'CRM login page is unavailable. Check the saved store URL and server connectivity.',
                'crm_export': 'CRM report export failed. Verify login, report selection, daily date controls and Excel export.',
                'wastage_report': 'The CRM did not retain the Wastage report selection. Check the report screen before retrying.',
                'wastage_filters': 'The CRM did not retain the requested Wastage day. Check Date Range, From Date and To Date before retrying.',
            }
            try:
                code = json.loads(downloaded.stderr).get('error_code')
            except (ValueError, TypeError, AttributeError):
                code = None
            # Never put arbitrary stderr or credential-bearing command arguments into audit logs.
            raise ValueError(safe_errors.get(code, 'Report downloader failed. Verify the saved CRM URL and server report configuration.'))
        if not output.is_file() or output.stat().st_size == 0:
            raise ValueError('Downloader did not produce the requested report')
        conn.execute("UPDATE data_import_runs SET status='downloaded', downloaded_at=now(), source_file_path=%s, source_start_date=%s, source_end_date=%s WHERE id=%s AND status='downloading'", (str(output), start, end, run_id))
        mapping_type = feed.get('mapping_type', 'store_product_code')
        mapping = feed['mapping']
        raw_rows = None
        verified_zero_activity = False
        if mapping.get('store'):
            # App parsers emit only code/date/quantity. Validate report identity
            # before normalization discards the original store column.
            raw_rows = read_rows(output, mapping)
            if any(str(row.get(mapping['store'], '')).strip() != identity[0] for row in raw_rows):
                raise ValueError('Report contains another store or an unrecognized store marker')
        if feed.get('normalize_with_app_parsers'):
            normalized = folder / 'normalized.json'
            app_dir = Path(config['app_directory']).resolve()
            normalized_result = subprocess.run([config.get('node_executable', 'node'), '--import', 'tsx',
                str(app_dir / 'automation/normalize-report.ts'), '--store-code', identity[0], '--store-name',
                str(conn.execute('SELECT name FROM stores WHERE id=%s', (store_id,)).fetchone()[0]),
                '--store-names', json.dumps([row[0] for row in conn.execute('SELECT name FROM stores').fetchall()]), '--input', str(output), '--output', str(normalized),
                '--source', source, '--start', start, '--end', end, '--erp', store.get('erp', 'msp')],
                cwd=app_dir, check=False, timeout=600, stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE, text=True)
            if normalized_result.returncode:
                # Only expose known, non-sensitive validation messages. Never publish
                # arbitrary subprocess stderr (paths, credentials or stack traces).
                diagnostic = (normalized_result.stderr or '').strip().splitlines()
                diagnostic = diagnostic[-1].strip() if diagnostic else ''
                safe_messages = (
                    'Report columns not detected',
                    'RTC report ',
                    'Empty file is not a recognized RTC report',
                    'Empty RTC report has no business date',
                    'Unreadable RTC workbook',
                    'Downloaded RTC file is a login page',
                    'Report date differs from requested date',
                    'Report contains another store',
                    'Wastage report must cover one valid day',
                    'Empty file is not a recognized Wastage report',
                    'Downloaded file is a login page',
                    'Report row date differs from requested date',
                    'Invalid requested dates',
                    'Sales report columns not detected',
                    'MSP Wastage report says No Records Available.',
                )
                if diagnostic.startswith(safe_messages):
                    raise ValueError('Report parsing failed: ' + diagnostic[:350] + ' Previous data retained.')
                raise ValueError('Report parsing failed. Verify the exported columns, daily date and ERP format; previous data retained.')
            envelope = json.loads(normalized.read_text(encoding='utf-8'))
            if envelope.get('source') != source:
                raise ValueError('Normalizer returned the wrong feed')
            review_descriptions = envelope.get('product_descriptions', {})
            verified_zero_activity = source in ('wastage', 'rtc') and envelope.get('zero_activity_verified') is True
            raw_rows = envelope['rows']
            mapping_type = envelope['mapping_type']
            mapping = {'code': 'Code', 'quantity': 'Qty', 'date': 'Date'}
        elif raw_rows is None:
            raw_rows = read_rows(output, mapping)
        prepared = prepare_rows(raw_rows, mapping, source, start, end, allow_empty=feed.get("allow_empty_report") is True or verified_zero_activity)
        stage_report(conn, (run_id, store_id, source, start, end), prepared, mapping_type, hashlib.sha256(output.read_bytes()).hexdigest())
        result = import_report(conn, (run_id, store_id, source, start, end), prepared, output.name, mapping_type, hashlib.sha256(output.read_bytes()).hexdigest())
        if result and result.get('details'):
            pending = result['details'].get('unmatched_barcodes', result['details'].get('unknown_product_codes', []))
            names = {code: str(review_descriptions[code])[:500] for code in pending if review_descriptions.get(code)}
            conn.execute("UPDATE data_import_runs SET details=details || %s::jsonb WHERE id=%s", (json.dumps({'unmatched_product_descriptions': names}), run_id))
        print(f'Job {run_id}: {result["status"]} ({result["saved"]} rows saved)', flush=True)
    except Exception as exc:
        # Keep credentials/command arguments out of logs and dashboard errors.
        if isinstance(exc, ImportValidationError) and isinstance(review_descriptions, dict):
            unknown = exc.details.get('unmatched_barcodes', exc.details.get('unknown_product_codes', []))
            exc.details['unmatched_product_descriptions'] = {
                code: str(review_descriptions[code])[:500] for code in unknown if review_descriptions.get(code)
            }
        message = str(exc) if isinstance(exc, ValueError) else type(exc).__name__
        failure_details = dict(exc.details if isinstance(exc, ImportValidationError) else {})
        failure_details.update({'failure_category': failure_category(message, exc), 'worker_revision': worker_revision()})
        conn.execute("UPDATE data_import_runs SET status='failed', completed_at=now(), error_message=%s, details=details || %s::jsonb WHERE id=%s AND status NOT IN ('success','failed')", (message[:2000], json.dumps(failure_details), run_id))
        print(f'Job {run_id}: failed ({message})', flush=True)


def require_reconciliation_schema(conn):
    ready = conn.execute("""SELECT to_regclass('crm_product_code_aliases') IS NOT NULL
        AND to_regclass('crm_report_stage') IS NOT NULL AND to_regclass('crm_product_identifiers') IS NOT NULL
        AND to_regclass('crm_partial_publications') IS NOT NULL""").fetchone()[0]
    if not ready:
        raise ValueError('Apply the backed-up 2026-10-20 reconciliation and partial-import migrations before starting this worker')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--config', required=True)
    parser.add_argument('--once', action='store_true')
    parser.add_argument('--max-jobs', type=int, help='Stop after this many queued jobs (successes or failures)')
    parser.add_argument('--store-id', type=int, help='Restrict a controlled database-backed run to one store')
    parser.add_argument('--retry-run', type=int, help='Create a new attempt for a failed run')
    parser.add_argument('--fail-stalled-minutes', type=int, help='Explicitly fail abandoned runs after stopping other workers')
    parser.add_argument('--enqueue-date', help='Queue all three movement feeds for YYYY-MM-DD')
    parser.add_argument('--enqueue-yesterday', action='store_true')
    parser.add_argument('--enqueue-only', action='store_true', help='Queue daily reports without consuming any jobs')
    args = parser.parse_args()
    if args.max_jobs is not None and args.max_jobs < 1:
        parser.error('--max-jobs must be positive')
    if args.store_id is not None and args.store_id < 1:
        parser.error('--store-id must be positive')
    if args.enqueue_only and not (args.enqueue_date or args.enqueue_yesterday):
        parser.error('--enqueue-only requires --enqueue-date or --enqueue-yesterday')
    config = json.loads(Path(args.config).read_text(encoding='utf-8-sig'))
    if args.store_id is not None:
        if config.get('store_settings_source') != 'database':
            parser.error('--store-id requires the database-backed profile')
        config['only_store_id'] = args.store_id
    import psycopg
    with psycopg.connect(os.environ['DATABASE_URL'], autocommit=True, prepare_threshold=None) as conn:
        require_reconciliation_schema(conn)
        if args.fail_stalled_minutes:
            if args.fail_stalled_minutes < 60:
                raise ValueError('Stalled timeout must be at least 60 minutes')
            conn.execute("UPDATE data_import_runs SET status='failed', completed_at=now(), error_message='Worker stopped or timed out' WHERE status IN ('downloading','downloaded','importing') AND started_at < now()-(%s * interval '1 minute')", (args.fail_stalled_minutes,))
        if args.retry_run:
            with conn.transaction():
                failed = conn.execute("SELECT store_id,source_type,source_start_date,source_end_date FROM data_import_runs WHERE id=%s AND status='failed' FOR UPDATE", (args.retry_run,)).fetchone()
                if not failed:
                    raise ValueError('Failed import run not found')
                store_id, source, start, end = failed
                conn.execute('SELECT pg_advisory_xact_lock(73501, %s)', (store_id,))
                active = conn.execute("SELECT id FROM data_import_runs WHERE store_id=%s AND source_type=%s AND status IN ('queued','downloading','downloaded','importing')", (store_id, source)).fetchone()
                if active:
                    raise ValueError('A run for this store/feed is already active')
                conn.execute("INSERT INTO data_import_runs(store_id,source_type,source_start_date,source_end_date,trigger_type,details) VALUES(%s,%s,%s,%s,'retry',%s::jsonb)", (store_id,source,start,end,json.dumps({'retry_of':args.retry_run})))
        if args.enqueue_date or args.enqueue_yesterday:
            day = args.enqueue_date or str(london_today() - timedelta(days=1))
            ids = nightly_store_ids(conn, config)
            if args.store_id is not None:
                ids = [store_id for store_id in ids if store_id == args.store_id]
            sources = enabled_movement_sources(config)
            enqueue_nightly(conn, ids, day, sources)
            if args.enqueue_yesterday:
                retried = enqueue_failed_retries(conn, ids, sources, day)
                print(f'Nightly recovery: {len(retried)} eligible failed reports queued for automatic retry.', flush=True)
        if args.enqueue_only:
            return
        processed = 0
        stopping = threading.Event()
        signal.signal(signal.SIGTERM, lambda *_: stopping.set())
        while not stopping.is_set():
            job = claim(conn, config)
            if job:
                process_job(conn, config, job)
                processed += 1
                if args.max_jobs is not None and processed >= args.max_jobs:
                    break
            elif args.once:
                break
            else:
                stopping.wait(10)


if __name__ == '__main__':
    main()
