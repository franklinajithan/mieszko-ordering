"""Offline report producer used in system tests; never opens a CRM session."""
import argparse
from openpyxl import Workbook
p=argparse.ArgumentParser()
for key in ('output','source','start','store'): p.add_argument('--'+key,required=True)
p.add_argument('--unknown',action='store_true')
p.add_argument('--empty',action='store_true')
p.add_argument('--report-store')
a=p.parse_args();book=Workbook();sheet=book.active
if a.source=='book_stock':
 sheet.append(['Item Code','Current Stock']);sheet.append(['10001',25]);sheet.append(['10002',0]);sheet.append(['10003',2])
elif a.source=='normal_sales':
 sheet.append(['Item Code','Sales Date','Sold Qty']);sheet.append(['10001',a.start,2*int(a.store)]);sheet.append(['10001',a.start,2*int(a.store)])
elif a.source=='rtc':
 sheet.append(['Item Code','RTC Qty']);sheet.append(['10001',int(a.store)])
else:
 sheet.append(['Item Code','Wastage Qty']);sheet.append(['10001',.5*int(a.store)])
if a.unknown: sheet.append(['UNKNOWN-PRODUCT',a.start,2] if a.source=='normal_sales' else ['UNKNOWN-PRODUCT',2])
if a.empty: sheet.delete_rows(2,sheet.max_row)
if a.report_store is not None:
 column=sheet.max_column+1
 sheet.cell(1,column,'Store Code')
 for row in range(2,sheet.max_row+1): sheet.cell(row,column,a.report_store)
book.save(a.output)
